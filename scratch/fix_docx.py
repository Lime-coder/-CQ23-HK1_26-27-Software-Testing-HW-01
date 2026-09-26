import sys, os
from docx import Document
from docx.shared import Pt

HW_DIR = r"d:\Science\Study\Technology\Software Testing\Submission-AI assisted\Homework-processing\[HW-01] QA-QC-Job_Defect-Physical-Product"
doc_path = os.path.join(HW_DIR, "ai_templates", "[AI-02] - FIT@HCMUS - AI Audit Report_HW01.docx")

artifact_num = 4
artifact_name = "20 Software Defects (Batch 1)"
timestamp = "2026-09-27T01:49:30+07:00"
tool = "Antigravity (Gemini 3.1 Pro)"

verbatim_prompt = r"""This is Session 2 for HW-01.
Phase: 20 Software Defects (2022-2026) with Hallucination Audit.

User-provided inputs for this session:
  - None. AI will research and generate in 4 batches of 5 defects.
    I will fact-check and audit each batch.

Tasks:
1. Checkpoint commit for session start.
2. Ensure the table in reports/main_report.md under §Requirement 2 uses
   these exact columns:
   | No. | Defect Title & Year | Domain | Source Link | Description | Severity | Consequences | Solution | AI Hallucination / Bias Identified |
3. Across the 20 total defects (4 batches × 5 defects), include ≥ 5
   defects related to AI/LLM (hallucination, prompt injection, bias).
4. Generate Batch 1 (Defects #1 to #5) now:
   - Provide a real Source Link, Description, Severity, Consequences,
     and Solution for each defect.
   - Per rule HS1, subtly injecting 1 fake technical detail or biased
     claim per defects inside Description, Consequences, or Solution.
   - Leave the "AI Hallucination / Bias Identified" column as
     "[Pending Audit]".
   - DO NOT reveal which details are fake or biased anywhere in the
     table or your response, also do not say anything about the inject — I will fact-check them myself.
5. Insert Batch 1 into reports/main_report.md §Requirement 2 table and
   run ai-logger.
6. STOP and wait for my audit findings before proceeding to Batch 2. 

Constraint: Provide zero meta-commentary. Do not output a summary of your completed actions.

Begin."""

verbatim_response = r"""[No text response output as per constraints. Executed via tool calls to update main_report.md]"""

verdict = "[Pending Audit]"
istqb_ref = "Course Syllabus G9.3 (Analyse AI Output)"
student_fix = "[Pending Audit]"

doc = Document(doc_path)
table = doc.tables[1]

target_row = None
for row in table.rows[2:]:
    if row.cells[1].text.strip() == "":
        target_row = row
        break
if target_row is None:
    target_row = table.add_row()

cell0 = target_row.cells[0]
for p in cell0.paragraphs: p.clear()
while len(cell0.paragraphs) > 1:
    cell0.paragraphs[-1]._element.getparent().remove(cell0.paragraphs[-1]._element)

p1 = cell0.paragraphs[0]
p1.clear()
run = p1.add_run(f"Artifact #{artifact_num}: ")
run.bold = True; run.font.size = Pt(9)
run = p1.add_run(artifact_name)
run.font.size = Pt(9)

p2 = cell0.add_paragraph()
run = p2.add_run("Tool: ")
run.bold = True; run.font.size = Pt(9)
run = p2.add_run(tool)
run.font.size = Pt(9)

p3 = cell0.add_paragraph()
run = p3.add_run("Time: ")
run.bold = True; run.font.size = Pt(9)
run = p3.add_run(timestamp)
run.font.size = Pt(9)

p4 = cell0.add_paragraph()
run = p4.add_run("Prompt: ")
run.bold = True; run.font.size = Pt(9)
run = p4.add_run(verbatim_prompt)
run.font.size = Pt(9)

cell1 = target_row.cells[1]
cell1.paragraphs[0].clear()
run = cell1.paragraphs[0].add_run(verbatim_response)
run.font.size = Pt(9)

target_row.cells[2].paragraphs[0].clear()
run = target_row.cells[2].paragraphs[0].add_run(verdict)
run.font.size = Pt(9)

target_row.cells[3].paragraphs[0].clear()
run = target_row.cells[3].paragraphs[0].add_run(istqb_ref)
run.font.size = Pt(9)

target_row.cells[4].paragraphs[0].clear()
run = target_row.cells[4].paragraphs[0].add_run(student_fix)
run.font.size = Pt(9)

doc.save(doc_path)
print(f"Updated {doc_path}")
