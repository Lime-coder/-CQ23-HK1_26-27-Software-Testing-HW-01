import sys, os, glob, re
import docx
from docx.shared import Pt

HW_DIR = r"d:\Science\Study\Technology\Software Testing\Submission-AI assisted\Homework-processing\[HW-01] QA-QC-Job_Defect-Physical-Product"

def get_fix_data(artifact_num):
    if artifact_num in [4, 6, 8]:
        return ("INVALID", "Course Syllabus G9.3 (Analyse) - Audit AI artifacts. Found hallucinated technical details.", "Identified fake details. Manually researched and provided corrections.")
    elif artifact_num in [5, 7, 9]:
        return ("VALID", "Course Syllabus G9.4 (Collaborate) - Pair AI + human.", "None - AI applied my corrections.")
    return None

# 1. Update DOCX
ai02_files = [os.path.join(HW_DIR, "ai_templates", f) for f in os.listdir(os.path.join(HW_DIR, "ai_templates")) if "AI-02" in f and not f.startswith("~")]
if not ai02_files:
    sys.exit("No docx found")
doc_path = ai02_files[0]
doc = docx.Document(doc_path)
table = doc.tables[1]

# Re-append Artifact #9 if missing
found_9 = False
for row in table.rows[2:]:
    if "Artifact #9" in row.cells[0].text:
        found_9 = True

if not found_9:
    target_row = table.add_row()
    artifact_num = 9
    artifact_name = "Defect Correction (Batch 4 Hallucinations)"
    timestamp = "2026-09-27T11:05:28+07:00"
    tool = "Antigravity (Gemini 3.1 Pro)"
    verbatim_prompt = r"""For Log4Shell, the fake detail is "exclusively affected Java versions strictly below 8u121, making newer systems immune by default." While newer Java versions mitigated one specific attack vector by default, Log4j itself was still highly vulnerable to other RCE chains and Denial of Service regardless of the Java version.  

For the Okta Support System Breach, the attackers didn't use an "SQL injection in the login.aspx page." They actually used a compromised service account and extracted stolen session tokens from HAR files that customers had uploaded to support tickets.   

For the Samsung Exynos Modem Baseband RCE, the flaw was triggered by malicious Session Description Protocol (SDP) packets sent over VoLTE or Wi-Fi calling using just the victim's phone number. It had nothing to do with sending a "malformed WhatsApp message"

For the Ivanti Connect Secure Zero-Day, the attackers didn't bounce connections through a "Tor network built into the appliance." The appliance doesn't have a built-in Tor network.

For the Tesla FSD "Rolling Stop" recall, the car rolling through stop signs was actually an intentional feature programmed into Tesla's "Assertive" driving profile. It was not because the AI "misclassified stop signs as speed limit signs due to a bug in the YOLOv4 object detection pipeline"."""

    verbatim_response = r"""[Tool calls executed silently to update reports/main_report.md with final Batch 4 corrections]"""
    
    cell0 = target_row.cells[0]
    for p in cell0.paragraphs: p.clear()
    while len(cell0.paragraphs) > 1:
        cell0.paragraphs[-1]._element.getparent().remove(cell0.paragraphs[-1]._element)
    
    p1 = cell0.paragraphs[0]
    run = p1.add_run(f"Artifact #{artifact_num}: ")
    run.bold = True
    run.font.size = Pt(9)
    run = p1.add_run(artifact_name)
    run.font.size = Pt(9)
    p2 = cell0.add_paragraph()
    run = p2.add_run("Tool: ")
    run.bold = True; run.font.size = Pt(9)
    run = p2.add_run(tool); run.font.size = Pt(9)
    p3 = cell0.add_paragraph()
    run = p3.add_run("Time: ")
    run.bold = True; run.font.size = Pt(9)
    run = p3.add_run(timestamp); run.font.size = Pt(9)
    p4 = cell0.add_paragraph()
    run = p4.add_run("Prompt: ")
    run.bold = True; run.font.size = Pt(9)
    run = p4.add_run(verbatim_prompt); run.font.size = Pt(9)
    
    cell1 = target_row.cells[1]
    cell1.paragraphs[0].clear()
    run = cell1.paragraphs[0].add_run(verbatim_response)
    run.font.size = Pt(9)
    
    target_row.cells[2].paragraphs[0].clear()
    target_row.cells[3].paragraphs[0].clear()
    target_row.cells[4].paragraphs[0].clear()

# Now go through all rows and fix Pending Audit
for row in table.rows[2:]:
    if len(row.cells) >= 5:
        text0 = row.cells[0].text
        num_match = re.search(r'Artifact #(\d+)', text0)
        if num_match:
            num = int(num_match.group(1))
            fix_data = get_fix_data(num)
            if fix_data:
                v, r, s = fix_data
                row.cells[2].paragraphs[0].text = v
                row.cells[3].paragraphs[0].text = r
                row.cells[4].paragraphs[0].text = s

doc.save(doc_path)
print("Docx fixed!")

# 2. Update prompt_log.md
with open(os.path.join(HW_DIR, "prompt_log.md"), "r", encoding="utf-8") as f:
    prompt_log = f.read()

for i in range(4, 10):
    fix_data = get_fix_data(i)
    if fix_data:
        v, r, s = fix_data
        block_start = prompt_log.find(f"### Artifact #{i}:")
        if block_start != -1:
            block_end = prompt_log.find(f"### Artifact #{i+1}:", block_start)
            if block_end == -1: block_end = len(prompt_log)
            block = prompt_log[block_start:block_end]
            
            block = re.sub(r'\*\*Student Audit Verdict:\*\* .*', f'**Student Audit Verdict:** {v}', block)
            block = re.sub(r'\*\*ISTQB / Slide Reference:\*\* .*', f'**ISTQB / Slide Reference:** {r}', block)
            block = re.sub(r'\*\*Student Modification:\*\* .*', f'**Student Modification:** {s}', block)
            
            prompt_log = prompt_log[:block_start] + block + prompt_log[block_end:]

with open(os.path.join(HW_DIR, "prompt_log.md"), "w", encoding="utf-8") as f:
    f.write(prompt_log)
print("prompt_log.md fixed!")

# 3. Update ai_audit_report.md
with open(os.path.join(HW_DIR, "reports", "ai_audit_report.md"), "r", encoding="utf-8") as f:
    audit_report = f.readlines()

for idx, line in enumerate(audit_report):
    num_match = re.search(r'\*\*Artifact #(\d+)\*\*', line)
    if num_match:
        num = int(num_match.group(1))
        fix_data = get_fix_data(num)
        if fix_data:
            v, r, s = fix_data
            parts = line.split(" | ")
            if len(parts) >= 5:
                parts[2] = f"**{v}**"
                parts[3] = r
                parts[4] = s + " |\n"
                audit_report[idx] = " | ".join(parts)

with open(os.path.join(HW_DIR, "reports", "ai_audit_report.md"), "w", encoding="utf-8") as f:
    f.writelines(audit_report)
print("ai_audit_report.md fixed!")
